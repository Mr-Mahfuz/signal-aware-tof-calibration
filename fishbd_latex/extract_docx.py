import docx
import os

docx_path = r'C:\Users\choto\OneDrive\Thesis\FISHBD3.docx'
out_dir = r'C:\Users\choto\OneDrive\Thesis\fishbd_latex'
out_file = os.path.join(out_dir, 'extracted_content.txt')

doc = docx.Document(docx_path)

with open(out_file, 'w', encoding='utf-8') as f:
    f.write("=== EXTRACTED TEXT ===\n\n")
    for para in doc.paragraphs:
        if para.text.strip():
            f.write(para.text + "\n\n")
            
    f.write("\n=== EXTRACTED TABLES ===\n\n")
    for i, table in enumerate(doc.tables):
        f.write(f"--- Table {i+1} ---\n")
        for row in table.rows:
            row_data = [cell.text.replace("\n", " ").strip() for cell in row.cells]
            f.write(" | ".join(row_data) + "\n")
        f.write("\n")

print(f"Extraction complete. Text and tables written to {out_file}")
